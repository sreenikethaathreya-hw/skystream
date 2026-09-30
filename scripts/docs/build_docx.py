"""Build docs/Defensible_Demand_Ledger.docx: the full application and data design document.

Run: MPLCONFIGDIR=/tmp/mpl python scripts/docs/build_docx.py  (needs python-docx, matplotlib, pillow)
"""

import json
import sys
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

sys.path.insert(0, str(Path(__file__).parent))
import content  # noqa: E402
import diagrams  # noqa: E402

ROOT = Path(__file__).resolve().parents[2]
SEED = ROOT / "data" / "seed"
SHOTS = ROOT / "tests" / "e2e" / "screenshots"
OUT_DIR = ROOT / "docs"
BRAND = RGBColor(0x1A, 0x5C, 0x31)


class Doc:
    def __init__(self) -> None:
        self.d = Document()
        section = self.d.sections[0]
        section.page_width, section.page_height = Cm(21.0), Cm(29.7)
        section.left_margin = section.right_margin = Cm(2.0)
        section.top_margin = section.bottom_margin = Cm(1.8)
        normal = self.d.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(10.5)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        for level, size in ((1, 18), (2, 14), (3, 12)):
            style = self.d.styles[f"Heading {level}"]
            style.font.name = "Calibri"
            style.font.size = Pt(size)
            style.font.color.rgb = BRAND
        self._figure = 0
        self._table = 0
        self._footer()

    def _footer(self) -> None:
        paragraph = self.d.sections[0].footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run("Defensible Demand Ledger - Application and Data Design  |  page ")
        run.font.size = Pt(8)
        self._field(paragraph, "PAGE", size=8)

    @staticmethod
    def _field(paragraph, code: str, size: int = 10) -> None:
        run = paragraph.add_run()
        run.font.size = Pt(size)
        begin, instr, end = OxmlElement("w:fldChar"), OxmlElement("w:instrText"), OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        instr.set(qn("xml:space"), "preserve")
        instr.text = code
        end.set(qn("w:fldCharType"), "end")
        run._r.append(begin)
        run._r.append(instr)
        run._r.append(end)

    def toc(self) -> None:
        paragraph = self.d.add_paragraph()
        self._field(paragraph, 'TOC \\o "1-3" \\h \\z \\u')
        note = self.d.add_paragraph("Right-click the table of contents and choose Update Field to refresh page numbers.")
        note.runs[0].italic = True
        note.runs[0].font.size = Pt(8.5)

    def h(self, text: str, level: int = 1) -> None:
        self.d.add_heading(text, level=level)

    def p(self, text: str, bold: bool = False, italic: bool = False, size: float | None = None) -> None:
        paragraph = self.d.add_paragraph()
        self._rich(paragraph, text, bold, italic, size)
        paragraph.paragraph_format.space_after = Pt(6)

    @staticmethod
    def _rich(paragraph, text: str, bold: bool = False, italic: bool = False, size: float | None = None) -> None:
        # **bold** segments inside text
        parts = text.split("**")
        for i, part in enumerate(parts):
            if not part:
                continue
            run = paragraph.add_run(part)
            run.bold = bold or i % 2 == 1
            run.italic = italic
            if size:
                run.font.size = Pt(size)

    def bullets(self, items: list[str], style: str = "List Bullet") -> None:
        for item in items:
            paragraph = self.d.add_paragraph(style=style)
            self._rich(paragraph, item)
            paragraph.paragraph_format.space_after = Pt(2)

    def numbered(self, items: list[str]) -> None:
        self.bullets(items, "List Number")

    def table(self, headers: list[str], rows: list[list], widths: list[float] | None = None, size: float = 8.5,
              caption: str | None = None) -> None:
        if caption:
            self._table += 1
            cap = self.d.add_paragraph()
            run = cap.add_run(f"Table {self._table}. {caption}")
            run.bold = True
            run.font.size = Pt(9)
            cap.paragraph_format.space_after = Pt(2)
        table = self.d.add_table(rows=1, cols=len(headers))
        table.style = "Light Grid Accent 1"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, header in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(header)
            run.bold = True
            run.font.size = Pt(size)
        for row in rows:
            cells = table.add_row().cells
            for i, value in enumerate(row):
                cells[i].text = ""
                self._rich(cells[i].paragraphs[0], str(value), size=size)
        if widths:
            for row in table.rows:
                for i, width in enumerate(widths):
                    row.cells[i].width = Cm(width)
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)

    def image(self, path: Path, caption: str, width_cm: float = 17.0) -> None:
        self._figure += 1
        self.d.add_picture(str(path), width=Cm(width_cm))
        self.d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = self.d.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cap.add_run(f"Figure {self._figure}. {caption}")
        run.italic = True
        run.font.size = Pt(9)

    def code(self, text: str) -> None:
        paragraph = self.d.add_paragraph()
        run = paragraph.add_run(text)
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
        shading = OxmlElement("w:shd")
        shading.set(qn("w:val"), "clear")
        shading.set(qn("w:fill"), "F4F5F2")
        paragraph.paragraph_format.element.get_or_add_pPr().append(shading)

    def page_break(self) -> None:
        self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def landscape_section(self, landscape: bool) -> None:
        section = self.d.add_section()
        if landscape:
            section.orientation = WD_ORIENT.LANDSCAPE
            section.page_width, section.page_height = Cm(29.7), Cm(21.0)
        else:
            section.orientation = WD_ORIENT.PORTRAIT
            section.page_width, section.page_height = Cm(21.0), Cm(29.7)

    def save(self, path: Path) -> None:
        self.d.save(str(path))


def seed(name: str):
    return json.loads((SEED / f"{name}.json").read_text())


def segment_rows() -> list[list[str]]:
    segments = seed("segments")
    market = {(m["segmentId"], m["year"]): m for m in seed("market_years")}
    plan = {(p["segmentId"], p["year"]): p for p in seed("plan_years")}
    owners = {"rep-a": "Rep A", "rep-b": "Rep B"}
    rows = []
    for s in segments:
        m, m25 = market[(s["id"], 2026)], market.get((s["id"], 2025))
        p = plan.get((s["id"], 2026))
        share = p["qtyKs"] / m["qtyKs"] if p and m["qtyKs"] else 0
        trend = (m["hectares"] - m25["hectares"]) / m25["hectares"] if m25 and m25["hectares"] else 0
        rows.append([
            s["id"], f"{s['profile'].replace('_', ' ').title()} {str(s['color']).title()}", owners[s["ownerId"]],
            f"{m['hectares']:,.0f}", f"{trend:+.1%}", f"{m['density']:.0f}", f"{m['qtyKs']:,.0f}",
            f"{m['priceExseed']:.0f}", f"{p['qtyKs']:,.0f}" if p else "-", f"{share:.1%}",
        ])
    return rows


def competitor_rows() -> list[list[str]]:
    rows = [c for c in seed("competitor_shares") if c["year"] == 2026]
    return [[c["competitor"], f"{c['sharePct']:.1f}%", f"EUR {c['valueEur'] / 1e6:.2f}M", c["trend"] or "-"]
            for c in sorted(rows, key=lambda c: -c["sharePct"])]


def crop_ledger() -> Path:
    src = SHOTS / "05-ledger.png"
    out = OUT_DIR / "diagrams" / "ledger-top.png"
    with Image.open(src) as img:
        img.crop((0, 0, img.width, min(img.height, 1150))).save(out)
    return out


def main() -> None:
    figures = diagrams.render_all(OUT_DIR / "diagrams")
    ledger_shot = crop_ledger()
    doc = Doc()
    content.build(doc, figures, SHOTS, ledger_shot, segment_rows(), competitor_rows(),
                  json.loads((SEED / "ingest_report.json").read_text()), date.today())
    out = OUT_DIR / "Defensible_Demand_Ledger.docx"
    doc.save(out)
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
